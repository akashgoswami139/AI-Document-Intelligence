from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import BinaryIO

import pandas as pd

from src.models import RawDocument, RawPage

logger = logging.getLogger(__name__)


class DocumentLoader:
    """Dispatches to format-specific loaders based on file extension."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx"}

    def load(self, file_path: str | Path, file_obj: BinaryIO | None = None) -> RawDocument:
        """Load a document from a file path or file-like object."""
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {ext}. "
                f"Supported: {', '.join(self.SUPPORTED_EXTENSIONS)}"
            )

        loader_map = {
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
            ".txt": self._load_text,
            ".md": self._load_text,
            ".csv": self._load_csv,
            ".xlsx": self._load_xlsx,
        }

        loader = loader_map[ext]
        logger.info("Loading %s with %s loader", path.name, ext)

        pages = loader(path, file_obj)

        return RawDocument(
            file_name=path.name,
            file_type=ext.lstrip("."),
            pages=pages,
            total_pages=len(pages),
        )

    # ── PDF ────────────────────────────────────────────────────────

    def _load_pdf(self, path: Path, file_obj: BinaryIO | None) -> list[RawPage]:
        from pypdf import PdfReader

        reader = PdfReader(file_obj if file_obj else str(path))
        pages = []

        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append(RawPage(
                page_number=i + 1,
                text=text.strip(),
                metadata={"char_count": len(text)},
            ))

        return pages

    # ── DOCX ───────────────────────────────────────────────────────

    def _load_docx(self, path: Path, file_obj: BinaryIO | None) -> list[RawPage]:
        from docx import Document as DocxDocument

        doc = DocxDocument(file_obj if file_obj else str(path))
        lines: list[str] = []
        tables_data: list[list[dict]] = []

        for para in doc.paragraphs:
            prefix = ""
            if para.style and para.style.name:
                style = para.style.name
                if style.startswith("Heading 1"):
                    prefix = "# "
                elif style.startswith("Heading 2"):
                    prefix = "## "
                elif style.startswith("Heading 3"):
                    prefix = "### "
                elif style.startswith("List"):
                    prefix = "- "
            lines.append(prefix + para.text)

        # Extract tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_data = {f"col_{j}": cell.text.strip() for j, cell in enumerate(row.cells)}
                table_rows.append(row_data)
            if table_rows:
                tables_data.append(table_rows)

        full_text = "\n".join(lines)

        # DOCX doesn't have native page numbers; treat as single page
        return [RawPage(
            page_number=1,
            text=full_text.strip(),
            tables=tables_data,
            metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
        )]

    # ── TXT / Markdown ─────────────────────────────────────────────

    def _load_text(self, path: Path, file_obj: BinaryIO | None) -> list[RawPage]:
        if file_obj:
            text = file_obj.read()
            if isinstance(text, bytes):
                text = text.decode("utf-8", errors="replace")
        else:
            text = path.read_text(encoding="utf-8", errors="replace")

        return [RawPage(
            page_number=1,
            text=text.strip(),
            metadata={"char_count": len(text)},
        )]

    # ── CSV ─────────────────────────────────────────────────────────

    def _load_csv(self, path: Path, file_obj: BinaryIO | None) -> list[RawPage]:
        if file_obj:
            content = file_obj.read()
            if isinstance(content, bytes):
                content = content.decode("utf-8", errors="replace")
            df = pd.read_csv(io.StringIO(content))
        else:
            df = pd.read_csv(path)

        # Convert to readable text with preserved column relationships
        text_lines = [f"Columns: {', '.join(df.columns.tolist())}"]
        text_lines.append(f"Total rows: {len(df)}\n")

        # Render as pipe-delimited table
        text_lines.append(df.to_markdown(index=False))

        # Also store as structured table data
        table_data = df.head(100).to_dict(orient="records")

        return [RawPage(
            page_number=1,
            text="\n".join(text_lines),
            tables=[table_data],
            metadata={"rows": len(df), "columns": len(df.columns)},
        )]

    # ── XLSX ────────────────────────────────────────────────────────

    def _load_xlsx(self, path: Path, file_obj: BinaryIO | None) -> list[RawPage]:
        source = file_obj if file_obj else str(path)
        xls = pd.ExcelFile(source, engine="openpyxl")
        pages = []

        for i, sheet_name in enumerate(xls.sheet_names):
            df = xls.parse(sheet_name)
            if df.empty:
                continue

            text_lines = [
                f"Sheet: {sheet_name}",
                f"Columns: {', '.join(str(c) for c in df.columns.tolist())}",
                f"Total rows: {len(df)}\n",
                df.to_markdown(index=False),
            ]

            table_data = df.head(100).to_dict(orient="records")

            pages.append(RawPage(
                page_number=i + 1,
                text="\n".join(text_lines),
                tables=[table_data],
                metadata={
                    "sheet_name": sheet_name,
                    "rows": len(df),
                    "columns": len(df.columns),
                },
            ))

        return pages if pages else [RawPage(page_number=1, text="[Empty spreadsheet]")]
